#!/usr/bin/env python3
"""Produce validated cover letters for an explicitly authorized Resume Ready queue.

This command is used at the end of an approved document-enabled search. It
only accepts the strict reviewer’s exact, dated Resume Ready queue and never
widens that queue or creates cover letters for diagnostic, applied, or
unreviewed postings.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import tempfile
import zipfile
from collections.abc import Callable, Mapping
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

if __package__:
    from scripts.docx_compatibility import (
        WORD_SETTINGS_PATH,
        modernize_docx_compatibility,
        require_modern_word_compatibility,
    )
    from scripts.mark_applied_jobs import ReportJob, load_tracker_jobs_snapshot, match_tracker_job
    from scripts.produce_resumes import REQUIRED_DOCX_PARTS, render_docx, validate_rendered_pages
    from scripts.resume_artifacts import ResumeReadyJob, artifact_slug, parse_resume_ready_jobs, run_date_from_ready_path
else:
    from docx_compatibility import (
        WORD_SETTINGS_PATH,
        modernize_docx_compatibility,
        require_modern_word_compatibility,
    )
    from mark_applied_jobs import ReportJob, load_tracker_jobs_snapshot, match_tracker_job
    from produce_resumes import REQUIRED_DOCX_PARTS, render_docx, validate_rendered_pages
    from resume_artifacts import ResumeReadyJob, artifact_slug, parse_resume_ready_jobs, run_date_from_ready_path


RenderFunction = Callable[[Path, Path], list[Path]]
LINKEDIN_URL_ENV = "JOBAGENT_LINKEDIN_URL"
PORTFOLIO_URL_ENV = "JOBAGENT_PORTFOLIO_URL"
SOURCE_RESUME_GLOB_ENV = "JOBAGENT_SOURCE_RESUME_GLOB"
SIGNATURE_TEXT_COLOR = RGBColor(89, 89, 89)
WRITING_VOICE = "candidate-confirmed-evidence"


@dataclass(frozen=True)
class Sender:
    name: str
    email: str
    phone: str
    location: str


@dataclass(frozen=True)
class CoverLetterProductionRecord:
    company: str
    title: str
    url: str
    template: str
    artifact_path: str
    status: str
    validation_status: str
    page_count: int
    artifact_sha256: str
    writing_voice: str


@dataclass(frozen=True)
class CoverLetterProductionResult:
    authorization_run_date: str
    generated_on: str
    records: list[CoverLetterProductionRecord]
    manifest_path: Path


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def human_date(iso_date: str) -> str:
    try:
        return dt.date.fromisoformat(iso_date).strftime("%B %-d, %Y")
    except ValueError as exc:
        raise ValueError(f"Cover-letter date must be YYYY-MM-DD: {iso_date}") from exc


def cover_letter_artifact_names(
    authorization_run_date: str,
    jobs: list[ResumeReadyJob],
    *,
    candidate_name: str = "Candidate Name",
) -> dict[ResumeReadyJob, str]:
    candidate_slug = artifact_slug(candidate_name)
    if not candidate_slug:
        raise ValueError("Cover-letter artifact names require a candidate name")
    base_names = [
        f"{authorization_run_date}-{artifact_slug(job.company)}-{artifact_slug(job.title)}-{candidate_slug}-Cover-Letter.docx"
        for job in jobs
    ]
    duplicate_bases = {name for name in base_names if base_names.count(name) > 1}
    names: dict[ResumeReadyJob, str] = {}
    for job, base_name in zip(jobs, base_names, strict=True):
        if base_name in duplicate_bases:
            url_hash = hashlib.sha256(job.url.encode("utf-8")).hexdigest()[:10]
            suffix = f"-posting-{url_hash}"
        else:
            suffix = ""
        names[job] = (
            f"{authorization_run_date}-{artifact_slug(job.company)}-{artifact_slug(job.title)}"
            f"{suffix}-{candidate_slug}-Cover-Letter.docx"
        )
    return names


def required_text(value: Any, *, field: str, source: Path) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{source} requires a non-empty {field!r} string")
    return value.strip()


def load_authorized_review(review_path: Path, *, authorization_run_date: str) -> Mapping[str, dict[str, Any]]:
    if not review_path.is_file():
        raise FileNotFoundError(f"Missing workflow review: {review_path}")
    payload: object = json.loads(review_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Expected a top-level object in workflow review: {review_path}")
    if required_text(payload.get("generated_on"), field="generated_on", source=review_path) != authorization_run_date:
        raise ValueError(f"Workflow review date does not match Resume Ready queue: {review_path}")
    raw_reviews = payload.get("reviews")
    if not isinstance(raw_reviews, list):
        raise ValueError(f"Workflow review requires a reviews list: {review_path}")
    reviews: dict[str, dict[str, Any]] = {}
    for item in raw_reviews:
        if not isinstance(item, dict):
            raise ValueError(f"Workflow review entries must be objects: {review_path}")
        url = required_text(item.get("url"), field="url", source=review_path)
        if url in reviews:
            raise ValueError(f"Workflow review contains duplicate posting URL: {url}")
        reviews[url] = item
    return reviews


def validate_authorization(
    *, jobs: list[ResumeReadyJob], review_path: Path, authorization_run_date: str, applied_jobs_path: Path
) -> None:
    reviews = load_authorized_review(review_path, authorization_run_date=authorization_run_date)
    tracker_jobs = load_tracker_jobs_snapshot(applied_jobs_path, expected_date=authorization_run_date)
    for job in jobs:
        review = reviews.get(job.url)
        if review is None:
            raise ValueError(f"Resume Ready posting is absent from strict review: {job.url}")
        if review.get("should_generate_resume") is not True:
            raise ValueError(f"Strict review does not authorize cover-letter production: {job.url}")
        if review.get("already_applied") is True:
            raise ValueError(f"Strict review marks posting already applied: {job.url}")
        matched = match_tracker_job(
            ReportJob(company=job.company, title=job.title, posted="", location="", fit=str(job.fit), url=job.url),
            tracker_jobs,
        )
        if matched is not None and matched.already_applied:
            raise ValueError(f"Resume Ready posting is already applied in jobtracker: {job.url}")


def load_sender(source_dir: Path) -> Sender:
    source_glob = os.environ.get(SOURCE_RESUME_GLOB_ENV, "*_SOURCE_*.docx")
    source_paths = sorted(source_dir.glob(source_glob))
    if not source_paths:
        raise FileNotFoundError(f"No source resume was found in {source_dir}")
    document = Document(source_paths[0])
    if not document.tables or not document.tables[0].rows or len(document.tables[0].columns) < 2:
        raise ValueError(f"Source resume has no usable sender header: {source_paths[0]}")
    left = [paragraph.text.strip() for paragraph in document.tables[0].cell(0, 0).paragraphs if paragraph.text.strip()]
    right = [paragraph.text.strip() for paragraph in document.tables[0].cell(0, 1).paragraphs if paragraph.text.strip()]
    if not left or not right:
        raise ValueError(f"Source resume sender header is incomplete: {source_paths[0]}")
    if len(right) != 2:
        raise ValueError(f"Source resume sender header must have identity and location lines: {source_paths[0]}")
    identity_line, location_line = right
    email_match = re.fullmatch(r"([^\s|]+@[^\s|]+)\s*•\s*LinkedIn\s*•\s*Portfolio", identity_line)
    if email_match is None:
        raise ValueError(f"Source resume sender identity line is invalid: {source_paths[0]}")
    phone_location = location_line.split(" • ", maxsplit=1)
    if len(phone_location) != 2 or not all(part.strip() for part in phone_location):
        raise ValueError(f"Source resume sender location line is invalid: {source_paths[0]}")
    return Sender(
        name=left[0],
        email=email_match.group(1),
        phone=phone_location[0].strip(),
        location=phone_location[1].strip(),
    )


def approved_voice_markers(job: ResumeReadyJob) -> tuple[str, str]:
    return (
        f"I am applying for the {job.title} position at {job.company}.",
        f"I would welcome a conversation about how that experience could support {job.company}.",
    )


def cover_letter_uses_approved_voice(path: Path, job: ResumeReadyJob) -> bool:
    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    return all(marker in text for marker in approved_voice_markers(job)) and (
        "I would welcome the opportunity to discuss" not in text
    )


def add_paragraph(document: Document, text: str, *, bold: bool = False, alignment: WD_ALIGN_PARAGRAPH | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    run.bold = bold


def add_external_hyperlink(paragraph: Paragraph, label: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_signature_run(run: Run, *, bold: bool = False, muted: bool = False) -> None:
    run.bold = bold
    if muted:
        run.font.size = Pt(9.5)
        run.font.color.rgb = SIGNATURE_TEXT_COLOR


def add_compact_paragraph(
    document: Document,
    text: str,
    *,
    bold: bool = False,
    muted: bool = False,
    space_after_points: float = 0,
) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after_points)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    style_signature_run(run, bold=bold, muted=muted)
    return paragraph


def sender_hyperlinks() -> dict[str, str]:
    return {
        "LinkedIn": os.environ.get(LINKEDIN_URL_ENV, "https://www.linkedin.com/in/candidate-profile"),
        "Portfolio": os.environ.get(PORTFOLIO_URL_ENV, "https://portfolio.example.com"),
    }


def add_sender_contact_line(document: Document, sender: Sender) -> None:
    paragraph = add_compact_paragraph(document, "")
    hyperlinks = sender_hyperlinks()
    labels = "|".join(re.escape(label) for label in hyperlinks)
    for segment in re.split(f"({labels})", "LinkedIn | Portfolio"):
        if not segment:
            continue
        target = hyperlinks.get(segment)
        if target is not None:
            add_external_hyperlink(paragraph, segment, target)
        else:
            run = paragraph.add_run(segment)
            style_signature_run(run, muted=True)


def add_sender_signature(document: Document, sender: Sender) -> None:
    add_compact_paragraph(document, "Sincerely,", space_after_points=8)
    add_compact_paragraph(document, sender.name, bold=True)
    add_compact_paragraph(document, sender.email, muted=True)
    add_sender_contact_line(document, sender)
    add_compact_paragraph(document, sender.phone, muted=True)
    add_compact_paragraph(document, sender.location, muted=True)


def build_cover_letter(*, destination_path: Path, sender: Sender, job: ResumeReadyJob, generated_on: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    add_paragraph(document, human_date(generated_on))
    add_paragraph(document, "Hiring Team")
    add_paragraph(document, job.company)
    add_paragraph(document, "Dear Hiring Team,")
    purpose_marker, closing_marker = approved_voice_markers(job)
    evidence_path = Path(os.environ.get("JOBAGENT_CANDIDATE_EVIDENCE", "candidate_evidence.json"))
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    statements = evidence.get("evidence")
    if not isinstance(statements, list) or not statements or any(not isinstance(item, str) or not item.strip() for item in statements):
        raise ValueError("Cover letters require confirmed statements in candidate_evidence.json")
    add_paragraph(document, purpose_marker)
    for statement in statements:
        add_paragraph(document, statement)
    add_paragraph(
        document,
        f"{closing_marker} Thank you for your consideration.",
    )
    add_sender_signature(document, sender)
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination_path)
    modernize_docx_compatibility(destination_path)


def validate_cover_letter(*, path: Path, job: ResumeReadyJob, generated_on: str) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError(f"Cover-letter artifact is missing or empty: {path}")
    if not zipfile.is_zipfile(path):
        raise ValueError(f"Cover-letter artifact is not a valid DOCX package: {path}")
    with zipfile.ZipFile(path) as archive:
        missing_parts = REQUIRED_DOCX_PARTS.difference(archive.namelist())
        if missing_parts:
            raise ValueError(f"Cover-letter artifact is missing DOCX parts: {sorted(missing_parts)}")
        require_modern_word_compatibility(
            archive.read(WORD_SETTINGS_PATH),
            document_path=path,
        )
    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    for expected in (job.company, job.title, human_date(generated_on)):
        if expected not in text:
            raise ValueError(f"Cover-letter artifact is missing expected text {expected!r}: {path}")


def write_manifest(
    path: Path,
    *,
    authorization_run_date: str,
    generated_on: str,
    records: list[CoverLetterProductionRecord],
) -> None:
    payload = {
        "authorization_run_date": authorization_run_date,
        "generated_on": generated_on,
        "writing_voice": WRITING_VOICE,
        "records": [asdict(record) for record in records],
    }
    temporary_path = path.with_suffix(".json.tmp")
    temporary_path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    temporary_path.replace(path)


def produce_cover_letters(
    *,
    resume_ready_path: Path,
    workflow_review_path: Path,
    applied_jobs_path: Path,
    source_dir: Path,
    output_dir: Path,
    generated_on: str,
    render_fn: RenderFunction | None = None,
) -> CoverLetterProductionResult:
    authorization_run_date = run_date_from_ready_path(resume_ready_path)
    jobs = parse_resume_ready_jobs(resume_ready_path)
    if not jobs:
        raise ValueError("Resume Ready queue is empty; no cover letters are authorized")
    validate_authorization(
        jobs=jobs,
        review_path=workflow_review_path,
        authorization_run_date=authorization_run_date,
        applied_jobs_path=applied_jobs_path,
    )
    sender = load_sender(source_dir)
    artifact_names = cover_letter_artifact_names(
        authorization_run_date,
        jobs,
        candidate_name=sender.name,
    )
    actual_render_fn = render_fn or render_docx
    records: list[CoverLetterProductionRecord] = []
    staged_paths: dict[str, Path] = {}
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="cover-letter-production-", dir=output_dir.parent) as temporary_dir:
        staging_root = Path(temporary_dir)
        for job in jobs:
            artifact_name = artifact_names[job]
            final_path = output_dir / artifact_name
            if final_path.exists() and cover_letter_uses_approved_voice(final_path, job):
                candidate_path = final_path
                status = "existing"
            else:
                candidate_path = staging_root / "documents" / artifact_name
                build_cover_letter(
                    destination_path=candidate_path,
                    sender=sender,
                    job=job,
                    generated_on=generated_on,
                )
                staged_paths[artifact_name] = candidate_path
                status = "generated"
            validate_cover_letter(path=candidate_path, job=job, generated_on=generated_on)
            page_paths = actual_render_fn(candidate_path, staging_root / "renders" / candidate_path.stem)
            validate_rendered_pages(page_paths, document_path=candidate_path)
            if len(page_paths) != 1:
                raise ValueError(f"Cover letter must fit on one page: {candidate_path}")
            records.append(
                CoverLetterProductionRecord(
                    company=job.company,
                    title=job.title,
                    url=job.url,
                    template=job.template,
                    artifact_path=str(final_path),
                    status=status,
                    validation_status="validated",
                    page_count=len(page_paths),
                    artifact_sha256=sha256_path(candidate_path),
                    writing_voice=WRITING_VOICE,
                )
            )
        output_dir.mkdir(parents=True, exist_ok=True)
        for artifact_name, staged_path in staged_paths.items():
            os.replace(staged_path, output_dir / artifact_name)
    manifest_path = resume_ready_path.parent / f"{authorization_run_date}-cover-letter-production.json"
    write_manifest(
        manifest_path,
        authorization_run_date=authorization_run_date,
        generated_on=generated_on,
        records=records,
    )
    return CoverLetterProductionResult(
        authorization_run_date=authorization_run_date,
        generated_on=generated_on,
        records=records,
        manifest_path=manifest_path,
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate validated DOCX cover letters only for strict Resume Ready postings."
    )
    parser.add_argument("--resume-ready", required=True, type=Path)
    parser.add_argument("--workflow-review", required=True, type=Path)
    parser.add_argument("--applied-jobs", required=True, type=Path)
    parser.add_argument("--source-dir", type=Path, default=Path("Source Docs"))
    parser.add_argument("--out", type=Path, default=Path("outputs/cover_letters"))
    parser.add_argument("--date", required=True, help="Human-facing letter date, YYYY-MM-DD")
    args = parser.parse_args()
    result = produce_cover_letters(
        resume_ready_path=args.resume_ready,
        workflow_review_path=args.workflow_review,
        applied_jobs_path=args.applied_jobs,
        source_dir=args.source_dir,
        output_dir=args.out,
        generated_on=args.date,
    )
    generated_count = sum(record.status == "generated" for record in result.records)
    existing_count = sum(record.status == "existing" for record in result.records)
    print(f"Validated {len(result.records)} cover letter(s): {generated_count} generated, {existing_count} existing")
    print(f"Wrote cover-letter production manifest: {result.manifest_path}")


if __name__ == "__main__":
    main()
